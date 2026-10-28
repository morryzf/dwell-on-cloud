import asyncio
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient

from app import auth, db, file_make
from app.main import FILE_TOOL_NAME, _chat_tools, _heartbeat_tool_allowed, app, file_tool


async def _run_chat_tool(chat_id, name, server, arguments):
    """cloud 的工具在聊天循环里直接调；这里照那一支的样子跑：出错变成给模型看的结果。"""
    assert (name, server) == (FILE_TOOL_NAME, "builtin:file")
    try:
        return await asyncio.to_thread(file_tool, chat_id, arguments), False
    except Exception as exc:
        return json.dumps({"is_error": True, "content": [{"type": "text", "text": str(exc)}]}, ensure_ascii=False), True

SAMPLE = """# 十月复习计划

这是一份**心理测量学**的复习计划。

## 第一周

- 信度
  - 重测信度
- 效度

1. 先看课本
2. 再做习题

> 慢慢来。

| 周次 | 内容 |
|---|---|
| 1 | 信度 |

---

```
alpha = k/(k-1)
```
"""


class MakeFileTest(unittest.TestCase):
    """文件名带扩展名决定格式；文字类原样存，Word / PDF 从 Markdown 排出来。"""

    def test_text_file_is_stored_as_is(self):
        name, mime, data = file_make.make_file("笔记.md", "# 标题\n内容")
        self.assertEqual((name, mime), ("笔记.md", "text/markdown"))
        self.assertEqual(data.decode("utf-8"), "# 标题\n内容")

    def test_csv_gets_bom_for_excel(self):
        _, _, data = file_make.make_file("清单.csv", "科目,时长\n信度,6\n")
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))

    def test_docx_has_headings_lists_and_table(self):
        import docx

        _, mime, data = file_make.make_file("计划.docx", SAMPLE)
        self.assertIn("wordprocessingml", mime)
        document = docx.Document(io.BytesIO(data))
        styles = [p.style.name for p in document.paragraphs]
        self.assertIn("Heading 1", styles)
        self.assertIn("List Bullet", styles)
        self.assertIn("List Bullet 2", styles)
        self.assertIn("List Number", styles)
        self.assertIn("Quote", styles)
        self.assertEqual(len(document.tables), 1)
        bold = [run.text for p in document.paragraphs for run in p.runs if run.bold]
        self.assertIn("心理测量学", bold)

    def test_xlsx_sheets_numbers_percent_formula_and_leading_zero(self):
        import openpyxl

        content = (
            "## 学习时长\n\n| 周次 | 时长 | 完成度 | 电话 |\n|---|---|---|---|\n"
            "| 1 | 6 | 80% | 013812345678 |\n| 2 | 8.5 | 12.5% | 1 |\n| 合计 | =SUM(B2:B3) | | |\n\n"
            "## 学习时长\n\n| a |\n|---|\n| x |\n"
        )
        _, mime, data = file_make.make_file("记录.xlsx", content)
        self.assertIn("spreadsheetml", mime)
        book = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ["学习时长", "学习时长 (2)"])
        sheet = book["学习时长"]
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual((sheet["A2"].value, sheet["B3"].value), (1, 8.5))
        self.assertEqual((sheet["C2"].value, sheet["C2"].number_format), (0.8, "0%"))
        self.assertEqual(sheet["C3"].number_format, "0.0#%")
        self.assertEqual(sheet["D2"].value, "013812345678")
        self.assertEqual(sheet["B4"].value, "=SUM(B2:B3)")

    def test_xlsx_from_csv_and_its_reading_view(self):
        import openpyxl

        _, _, data = file_make.make_file("分数.xlsx", "名字,分数\n小云,90\n小雨,85\n")
        rows = [[c.value for c in r] for r in openpyxl.load_workbook(io.BytesIO(data)).active.iter_rows()]
        self.assertEqual(rows, [["名字", "分数"], ["小云", 90], ["小雨", 85]])
        page = file_make.xlsx_to_html(data, "分数.xlsx", "../x?download=1")
        self.assertIn("<th>名字</th>", page)
        self.assertIn('<td class="num">90</td>', page)
        self.assertIn("下载 Excel", page)

    def test_reading_view_back_button_is_the_only_script_allowed(self):
        import base64
        import hashlib
        import re as regex

        page = file_make.docx_to_html(file_make.make_file("a.docx", "字")[2], "a.docx", "../x?download=1")
        self.assertIn('id="back"', page)
        self.assertIn('id="dl"', page)
        self.assertIn('id="dlnote" class="dlnote" hidden', page)
        scripts = regex.findall(r"<script>(.*?)</script>", page, regex.S)
        self.assertEqual(len(scripts), 1)
        digest = base64.b64encode(hashlib.sha256(scripts[0].encode()).digest()).decode()
        self.assertIn("'sha256-" + digest + "'", file_make.PREVIEW_CSP)
        self.assertIn("default-src 'none'", file_make.PREVIEW_CSP)

    def test_pdf_embeds_chinese_text(self):
        if not file_make.find_cjk_font()[0]:
            self.skipTest("这台机器没有中文字体")
        from pypdf import PdfReader

        _, mime, data = file_make.make_file("计划.pdf", SAMPLE)
        self.assertEqual(mime, "application/pdf")
        text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
        self.assertIn("十月复习计划", text.replace("​", ""))

    def test_pdf_without_font_says_so(self):
        original = file_make.find_cjk_font, file_make.ensure_cjk_font
        file_make.find_cjk_font = lambda: ("", "")
        file_make.ensure_cjk_font = lambda timeout=90: ("", "")
        try:
            with self.assertRaisesRegex(ValueError, "中文字体"):
                file_make.make_file("计划.pdf", "内容")
        finally:
            file_make.find_cjk_font, file_make.ensure_cjk_font = original

    def test_missing_font_is_downloaded_once_and_bad_downloads_are_rejected(self):
        """系统没字体就下载到 data/fonts；第一个镜像给了个网页（不是字体）就换下一个。"""
        calls = []

        class Fake(io.BytesIO):
            def __enter__(self): return self
            def __exit__(self, *exc): return False

        def fake_urlopen(url, timeout=0):
            calls.append(url)
            if len(calls) == 1:
                return Fake(b"<html>blocked</html>")
            return Fake(b"OTTO" + b"\0" * 1_200_000)

        with tempfile.TemporaryDirectory() as tmp:
            saved = (file_make.FONT_DIR, list(file_make._FONT_CANDIDATES), file_make.urllib.request.urlopen)
            file_make.FONT_DIR = Path(tmp)
            file_make._FONT_CANDIDATES[:] = []
            file_make.urllib.request.urlopen = fake_urlopen
            try:
                import os
                os.environ.pop("DWELL_PDF_FONT", None)
                regular, bold = file_make.ensure_cjk_font()
                self.assertTrue(regular.endswith("NotoSansSC-Regular.otf"))
                self.assertTrue(bold.endswith("NotoSansSC-Bold.otf"))
                self.assertEqual(len(calls), 3)              # 第一个镜像不算数，换了一个；粗体一次成
                self.assertEqual(file_make.ensure_cjk_font(), (regular, bold))
                self.assertEqual(len(calls), 3)              # 下过了就不再下
            finally:
                file_make.FONT_DIR, candidates, file_make.urllib.request.urlopen = saved
                file_make._FONT_CANDIDATES[:] = candidates

    def test_word_reading_view_keeps_structure_and_escapes(self):
        _, _, data = file_make.make_file("计划.docx", SAMPLE + "\n\n<script>alert(1)</script>\n")
        page = file_make.docx_to_html(data, "计划.docx", "../x?download=1")
        for piece in ("<h1>十月复习计划</h1>", "<b>心理测量学</b>", "<ul>", "<ol>",
                      "<blockquote>", "<table>", "<pre>", "<hr>", 'href="../x?download=1" target="_blank"'):
            self.assertIn(piece, page)
        self.assertNotIn("<script>alert", page)        # 模型写的脚本被转义；页面里只有返回键那一段
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", page)

    def test_bad_names_and_content(self):
        for name in ("没有扩展名", "病毒.exe"):
            with self.assertRaises(ValueError):
                file_make.make_file(name, "内容")
        with self.assertRaises(ValueError):
            file_make.make_file("空.txt", "   ")
        self.assertEqual(file_make.clean_name("../../etc/密码.txt"), "密码.txt")

    def test_chinese_gets_break_points_but_not_before_closing_punctuation(self):
        text = file_make.breakable("中文，好。abc")
        self.assertNotIn("​，", text)
        self.assertNotIn("​。", text)
        self.assertIn("中​文", text)
        self.assertNotIn("a​b", text)


class FileToolFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.previous_path = db.DB_PATH
        db.DB_PATH = str(Path(cls.tmp.name) / "files.db")
        db.init_db()
        app.dependency_overrides[auth.require_auth] = lambda: None
        cls.client = TestClient(app)
        cls.chat = db.chat_add("文件")

    @classmethod
    def tearDownClass(cls):
        app.dependency_overrides.pop(auth.require_auth, None)
        db.DB_PATH = cls.previous_path
        cls.tmp.cleanup()

    def test_tool_is_offered_in_every_chat_but_not_to_heartbeat(self):
        tools, tool_map = asyncio.run(_chat_tools(self.chat["id"]))
        names = [tool["function"]["name"] for tool in tools]
        self.assertIn(FILE_TOOL_NAME, names)
        tool = next(tool for tool in tools if tool["function"]["name"] == FILE_TOOL_NAME)
        self.assertFalse(_heartbeat_tool_allowed(tool, tool_map[FILE_TOOL_NAME]))

    def test_created_file_downloads_with_its_name(self):
        result, is_error = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file", {"filename": "笔记.txt", "content": "你好"}))
        self.assertFalse(is_error)
        file = json.loads(result)["file"]
        response = self.client.get("/api/files/" + file["id"])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode("utf-8"), "你好")
        self.assertIn("charset=utf-8", response.headers["content-type"])
        self.assertIn("inline", response.headers["content-disposition"])
        self.assertIn("filename*=UTF-8''%E7%AC%94%E8%AE%B0.txt", response.headers["content-disposition"])
        saved = self.client.get("/api/files/" + file["id"] + "?download=1")
        self.assertIn("attachment", saved.headers["content-disposition"])

    def test_html_is_sandboxed(self):
        result, _ = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file",
            {"filename": "页.html", "content": "<script>alert(1)</script>"}))
        file = json.loads(result)["file"]
        response = self.client.get("/api/files/" + file["id"])
        self.assertEqual(response.headers.get("content-security-policy"), "sandbox")

    def test_tool_error_goes_back_to_the_model(self):
        result, is_error = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file", {"filename": "x.exe", "content": "a"}))
        self.assertTrue(is_error)
        self.assertIn("不支持", result)

    def test_missing_file_is_404(self):
        self.assertEqual(self.client.get("/api/files/nope").status_code, 404)
        self.assertEqual(self.client.get("/api/files/nope/view").status_code, 404)

    def test_word_opens_in_reading_view_others_go_straight_to_file(self):
        result, _ = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file", {"filename": "计划.docx", "content": SAMPLE}))
        docx_id = json.loads(result)["file"]["id"]
        page = self.client.get("/api/files/" + docx_id + "/view")
        self.assertEqual(page.status_code, 200)
        self.assertIn("text/html", page.headers["content-type"])
        self.assertEqual(page.headers["content-security-policy"], file_make.PREVIEW_CSP)
        self.assertIn("十月复习计划", page.text)
        result, _ = asyncio.run(_run_chat_tool(
            self.chat["id"], FILE_TOOL_NAME, "builtin:file", {"filename": "a.txt", "content": "x"}))
        txt_id = json.loads(result)["file"]["id"]
        moved = self.client.get("/api/files/" + txt_id + "/view", follow_redirects=False)
        self.assertEqual(moved.status_code, 307)


class CloudChatLoopTest(unittest.TestCase):
    def test_chat_loop_runs_the_file_tool_in_a_thread_and_slims_the_record(self):
        main = (Path(__file__).resolve().parents[1] / "app" / "main.py").read_text(encoding="utf-8")
        branch = main[main.index('elif server == "builtin:file":'):][:400]
        self.assertIn("await asyncio.to_thread(file_tool, chat_id, arguments)", branch)
        self.assertIn('_slim_file_tool_call(record["id"], name, arguments, is_error)', branch)


class FileCardStaticTest(unittest.TestCase):
    def test_page_draws_a_card_for_created_files(self):
        html = (Path(__file__).resolve().parents[1] / "static" / "index.html").read_text(encoding="utf-8")
        self.assertIn("if (t.name === 'DwellFileCreate' && !isError) addGeneratedFileCard(t);", html)
        self.assertIn("save.href = 'api/files/' + encodeURIComponent(file.id) + '?download=1';", html)
        self.assertIn("(ext === 'DOCX' || ext === 'XLSX' ? '/view' : '')", html)


if __name__ == "__main__":
    unittest.main()
