"""模型在聊天里生成的文件：文字类原样存，Word / PDF 由这里排出来。

模型交来的是「文件名 + 内容」。文字类（txt、md、csv、代码……）内容就是文件本身；
docx 和 pdf 的内容按 Markdown 写——标题、段落、列表、粗体斜体、引用、代码块、
分隔线、表格——这里照着排成真正的文档。

PDF 要嵌一个带中文的字体，不然中文全是方块。字体不进仓库（十几 MB）：先用服务器
系统里装着的（Noto CJK、文泉驿……）；没有就把思源黑体（Noto Sans SC）下载到
data/fonts 里，以后一直用这份。都不行就报一句清楚的错，让模型换成 Word。
可以用环境变量 DWELL_PDF_FONT 直接指一个字体文件。

Excel（.xlsx）的内容写成一张或几张 Markdown 表格（表格上面的标题就是工作表名），
或者直接写 CSV。数字存成数字、百分数存成百分比、「=」开头的是公式。

Word 和 Excel 在 iPhone 的浏览器里都打不开（一片空白），所以另给一个「阅读视图」：
docx_to_html / xlsx_to_html 把做好的文件读回来，变成一页能看的网页。
"""
from __future__ import annotations

import base64
import csv
import hashlib
import html
import io
import os
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

MAX_CONTENT_CHARS = 400_000
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_NAME_CHARS = 120

TEXT_TYPES = {
    "txt": "text/plain", "md": "text/markdown", "markdown": "text/markdown",
    "csv": "text/csv", "tsv": "text/tab-separated-values", "json": "application/json",
    "html": "text/html", "htm": "text/html", "xml": "application/xml",
    "yaml": "text/yaml", "yml": "text/yaml", "toml": "text/plain", "ini": "text/plain",
    "log": "text/plain", "srt": "text/plain", "vtt": "text/vtt", "tex": "text/plain",
    "py": "text/plain", "js": "text/plain", "ts": "text/plain", "css": "text/css",
    "sh": "text/plain", "sql": "text/plain", "swift": "text/plain", "java": "text/plain",
    "c": "text/plain", "cpp": "text/plain", "h": "text/plain", "go": "text/plain",
    "rs": "text/plain", "rb": "text/plain", "php": "text/plain", "kt": "text/plain",
}
DOC_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
PREVIEW_TYPES = {"docx", "xlsx"}       # 这两种点卡片进阅读视图
SUPPORTED = sorted(TEXT_TYPES) + sorted(DOC_TYPES)


def clean_name(raw: object) -> str:
    """文件名只留能安全落地的字：去掉路径、控制字符，太长就截。"""
    name = unicodedata.normalize("NFC", str(raw or "")).strip()
    name = name.replace("\\", "/").split("/")[-1]
    name = "".join(ch for ch in name if ch.isprintable() and ch not in '<>:"|?*')
    name = re.sub(r"\s+", " ", name).strip(" .")
    if not name:
        raise ValueError("要有文件名，带扩展名，比如 笔记.md")
    if "." not in name:
        raise ValueError("文件名要带扩展名，比如 笔记.md、报告.docx、清单.pdf")
    stem, ext = name.rsplit(".", 1)
    if len(name) > MAX_NAME_CHARS:
        name = stem[:MAX_NAME_CHARS - len(ext) - 1] + "." + ext
    return name


def make_file(filename: object, content: object) -> tuple[str, str, bytes]:
    """返回 (文件名, MIME, 文件内容)。"""
    name = clean_name(filename)
    ext = name.rsplit(".", 1)[1].lower()
    text = str(content or "")
    if not text.strip():
        raise ValueError("文件内容是空的")
    if len(text) > MAX_CONTENT_CHARS:
        raise ValueError(f"内容太长了（{len(text)} 字），一份文件最多 {MAX_CONTENT_CHARS} 字")
    if ext in TEXT_TYPES:
        data = text.encode("utf-8")
        if ext in {"csv", "tsv"}:
            data = b"\xef\xbb\xbf" + data      # 带 BOM，Excel 打开中文才不乱码
        mime = TEXT_TYPES[ext]
    elif ext == "docx":
        data, mime = markdown_to_docx(text), DOC_TYPES["docx"]
    elif ext == "xlsx":
        data, mime = content_to_xlsx(text), DOC_TYPES["xlsx"]
    elif ext == "pdf":
        data, mime = markdown_to_pdf(text), DOC_TYPES["pdf"]
    else:
        raise ValueError("不支持 ." + ext + "。能做的：" + "、".join(SUPPORTED))
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("生成的文件超过 10 MB，拆成几份再做")
    return name, mime, data


# ---------------------------------------------------------------- Markdown 拆块

_LIST = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def blocks(text: str) -> list[dict]:
    """把 Markdown 拆成块：heading / para / list / quote / code / rule / table。"""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    out: list[dict] = []
    para: list[str] = []
    i = 0

    def flush():
        if para:
            out.append({"type": "para", "text": " ".join(s.strip() for s in para)})
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            out.append({"type": "code", "text": "\n".join(body)})
            i += 1
            continue
        if not stripped:
            flush()
            i += 1
            continue
        m = _HEADING.match(stripped)
        if m:
            flush()
            out.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2)})
            i += 1
            continue
        if _RULE.match(stripped):
            flush()
            out.append({"type": "rule"})
            i += 1
            continue
        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_cells(stripped)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            out.append({"type": "table", "rows": [r + [""] * (width - len(r)) for r in rows]})
            continue
        m = _LIST.match(line)
        if m:
            flush()
            items = []
            while i < len(lines):
                m = _LIST.match(lines[i])
                if not m:
                    break
                items.append({"level": len(m.group(1).expandtabs(4)) // 2,
                              "ordered": m.group(2)[0].isdigit(), "text": m.group(3)})
                i += 1
            out.append({"type": "list", "items": items})
            continue
        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            out.append({"type": "quote", "text": " ".join(q for q in quote if q)})
            continue
        para.append(line)
        i += 1
    flush()
    return out


_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`)")


def inline_runs(text: str) -> list[tuple[str, str]]:
    """一行字拆成 (样式, 字)：样式是 '' / 'b' / 'i' / 'code'。链接只留文字。"""
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and len(part) > 4:
            runs.append(("b", part[2:-2]))
        elif part.startswith("`") and len(part) > 2:
            runs.append(("code", part[1:-1]))
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            runs.append(("i", part[1:-1]))
        else:
            runs.append(("", part))
    return runs


def plain(text: str) -> str:
    return "".join(t for _, t in inline_runs(text))


# ---------------------------------------------------------------- Word

def markdown_to_docx(text: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
                       "Heading 4", "Quote", "List Bullet", "List Number"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Helvetica"
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        fonts.set(qn("w:eastAsia"), "PingFang SC")
        if style_name.startswith("Heading") or style_name == "Title":
            style.font.color.rgb = RGBColor(0x2B, 0x26, 0x29)
    doc.styles["Normal"].font.size = Pt(11)

    def add_runs(paragraph, line: str):
        for kind, chunk in inline_runs(line):
            run = paragraph.add_run(chunk)
            if kind == "b":
                run.bold = True
            elif kind == "i":
                run.italic = True
            elif kind == "code":
                run.font.name = "Menlo"
                run.font.size = Pt(10)

    for block in blocks(text):
        kind = block["type"]
        if kind == "heading":
            add_runs(doc.add_heading("", level=min(block["level"], 4)), block["text"])
        elif kind == "para":
            add_runs(doc.add_paragraph(), block["text"])
        elif kind == "list":
            for item in block["items"]:
                style = "List Number" if item["ordered"] else "List Bullet"
                if item["level"] >= 1:
                    style += " 2"
                try:
                    paragraph = doc.add_paragraph(style=style)
                except KeyError:
                    paragraph = doc.add_paragraph(style="List Number" if item["ordered"] else "List Bullet")
                add_runs(paragraph, item["text"])
        elif kind == "quote":
            add_runs(doc.add_paragraph(style="Quote"), block["text"])
        elif kind == "code":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.name = "Menlo"
            run.font.size = Pt(9.5)
        elif kind == "rule":
            paragraph = doc.add_paragraph("—" * 12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "table":
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    add_runs(paragraph, cell)
                    if r == 0:
                        for run in paragraph.runs:
                            run.bold = True
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- PDF

_FONT_CANDIDATES = [
    # (常规, 粗体)；粗体没有就用常规的
    ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ("/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc", "/usr/share/fonts/noto-cjk/NotoSansCJK-Bold.ttc"),
    ("/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc", "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Bold.ttc"),
    ("/usr/share/fonts/google-noto-sans-cjk-fonts/NotoSansCJK-Regular.ttc", "/usr/share/fonts/google-noto-sans-cjk-fonts/NotoSansCJK-Bold.ttc"),
    ("/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc"),
    ("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", ""),
    ("/usr/share/fonts/wqy-microhei/wqy-microhei.ttc", ""),
    ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", ""),
    ("/usr/share/fonts/wqy-zenhei/wqy-zenhei.ttc", ""),
    ("/System/Library/Fonts/PingFang.ttc", ""),
]


FONT_DIR = Path(os.environ.get("DWELL_FONT_DIR", "./data/fonts"))
_FONT_FILES = ("NotoSansSC-Regular.otf", "NotoSansSC-Bold.otf")
# 国内服务器上哪个通不好说，挨个试。都是 Google 开源的思源黑体（OFL 授权）。
_FONT_MIRRORS = (
    "https://cdn.jsdelivr.net/gh/notofonts/noto-cjk@main/Sans/SubsetOTF/SC/{name}",
    "https://fastly.jsdelivr.net/gh/notofonts/noto-cjk@main/Sans/SubsetOTF/SC/{name}",
    "https://gcore.jsdelivr.net/gh/notofonts/noto-cjk@main/Sans/SubsetOTF/SC/{name}",
    "https://raw.githubusercontent.com/notofonts/noto-cjk/main/Sans/SubsetOTF/SC/{name}",
)


def _downloaded_font(name: str) -> str:
    path = FONT_DIR / name
    return str(path) if path.is_file() and path.stat().st_size > 1_000_000 else ""


def find_cjk_font() -> tuple[str, str]:
    """(常规字体路径, 粗体字体路径)。找不到返回 ('', '')。"""
    forced = os.environ.get("DWELL_PDF_FONT", "").strip()
    if forced and os.path.isfile(forced):
        return forced, ""
    for regular, bold in _FONT_CANDIDATES:
        if os.path.isfile(regular):
            return regular, bold if bold and os.path.isfile(bold) else ""
    regular = _downloaded_font(_FONT_FILES[0])
    if regular:
        return regular, _downloaded_font(_FONT_FILES[1])
    return "", ""


def ensure_cjk_font(timeout: float = 90) -> tuple[str, str]:
    """系统里没有中文字体就下载一份到 data/fonts。返回 find_cjk_font() 的结果。"""
    found = find_cjk_font()
    if found[0]:
        return found
    FONT_DIR.mkdir(parents=True, exist_ok=True)
    for name in _FONT_FILES:
        if _downloaded_font(name):
            continue
        for mirror in _FONT_MIRRORS:
            url = mirror.format(name=name)
            part = FONT_DIR / (name + ".part")
            try:
                with urllib.request.urlopen(url, timeout=timeout) as response, open(part, "wb") as out:
                    while chunk := response.read(1 << 16):
                        out.write(chunk)
                with open(part, "rb") as check:
                    head = check.read(4)
                if head != b"OTTO" or part.stat().st_size < 1_000_000:
                    raise ValueError("下下来的不是字体文件")
                part.replace(FONT_DIR / name)
                break
            except Exception:
                part.unlink(missing_ok=True)
    return find_cjk_font()


def _collection_index(path: str) -> int:
    """.ttc 里有好几套字体（日韩简繁）。挑简体中文那套，挑不出就第一套。"""
    if not path.lower().endswith(".ttc"):
        return 0
    try:
        from fontTools.ttLib import TTCollection
        names = []
        for font in TTCollection(path).fonts:
            names.append(str(font["name"].getDebugName(1) or ""))
        for wanted in (" SC", "SC ", "CN", "GB", "Hei", "Micro"):
            for index, name in enumerate(names):
                if wanted in name:
                    return index
    except Exception:
        pass
    return 0


_CLOSE = set("，。、；：？！）》」』】〉”’…—,.;:?!)]}%")
_OPEN = set("（《「『【〈“‘([{")


def _cjk(ch: str) -> bool:
    code = ord(ch)
    return (0x2E80 <= code <= 0x9FFF or 0xF900 <= code <= 0xFAFF
            or 0xFF00 <= code <= 0xFFEF or 0x3000 <= code <= 0x303F)


def breakable(text: str) -> str:
    """PDF 那边只在空格处换行，中文没有空格，一整段会被当成一个「词」整块挪到下一行。
    在中文字之间塞零宽空格当换行点；不塞在「，。」这类收尾标点前面，免得它跑到行首。"""
    out = []
    for i, ch in enumerate(text):
        out.append(ch)
        nxt = text[i + 1] if i + 1 < len(text) else ""
        if nxt and (_cjk(ch) or _cjk(nxt)) and nxt not in _CLOSE and ch not in _OPEN \
                and not ch.isspace() and not nxt.isspace():
            out.append("\u200b")
    return "".join(out)


def markdown_to_pdf(text: str) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import TextMode, XPos, YPos

    regular, bold = find_cjk_font()
    if not regular:
        regular, bold = ensure_cjk_font()
    if not regular:
        raise ValueError("服务器上没有中文字体，也没能下载到，PDF 排不出来。先做成 .docx，或者重新部署一次再试")

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("cjk", "", regular, collection_font_number=_collection_index(regular))
    bold_path = bold or regular
    pdf.add_font("cjk", "B", bold_path, collection_font_number=_collection_index(bold_path))
    # 零宽空格只是换行点。字体里没这个字形时库会按默认宽度往前挪，字却不画，
    # 每段粗体、代码的前后就空出一截——把它的宽度钉成 0。
    for font in pdf.fonts.values():
        if hasattr(font, "cw"):
            font.cw[0x200B] = 0
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    ink, dim = (43, 38, 41), (120, 112, 116)
    pdf.set_text_color(*ink)

    def run_font(bold_run: bool, size: float):
        """有粗体字体就用；没有（文泉驿这种只有一个字重）就描边假装加粗。"""
        pdf.set_font("cjk", "B" if bold_run else "", size)
        if bold_run and not bold:
            pdf.text_mode = TextMode.FILL_STROKE
            pdf.set_line_width(size * 0.018)
            pdf.set_draw_color(*pdf_ink())
        else:
            pdf.text_mode = TextMode.FILL

    def pdf_ink():
        return dim if quoting[0] else ink

    quoting = [False]

    def write_runs(line: str, size: float, height: float, bold_all: bool = False):
        for kind, chunk in inline_runs(line):
            run_font(bold_all or kind == "b", size)
            pdf.write(height, breakable(chunk))
        pdf.text_mode = TextMode.FILL

    def write_line(line: str, size: float, height: float, indent: float = 0, bold_all: bool = False):
        """一段带粗体的字，自动换行。斜体和行内代码在 PDF 里按普通字排。"""
        pdf.set_left_margin(pdf.l_margin + indent)
        pdf.set_x(pdf.l_margin)
        write_runs(line, size, height, bold_all)
        pdf.set_left_margin(pdf.l_margin - indent)
        pdf.ln(height)

    heading_sizes = {1: 20, 2: 16, 3: 14, 4: 12.5, 5: 12, 6: 11.5}
    for block in blocks(text):
        kind = block["type"]
        if kind == "heading":
            size = heading_sizes[block["level"]]
            pdf.ln(3 if block["level"] > 1 else 1)
            write_line(block["text"], size, size * 0.55, bold_all=True)
            pdf.ln(1.5)
        elif kind == "para":
            write_line(block["text"], 11, 6.4)
            pdf.ln(2.2)
        elif kind == "list":
            number = 0
            for item in block["items"]:
                indent = 4 + item["level"] * 6
                if item["ordered"]:
                    number += 1
                    mark = f"{number}."
                else:
                    mark = "●"                # 「•」不少中文字体里没有，「●」都有
                pdf.set_font("cjk", "", 6.5 if mark == "●" else 11)
                pdf.set_x(pdf.l_margin + indent)
                pdf.cell(6, 6.4, mark)
                pdf.set_left_margin(pdf.l_margin + indent + 6)
                write_runs(item["text"], 11, 6.4)
                pdf.set_left_margin(pdf.l_margin - indent - 6)
                pdf.ln(6.4)
            pdf.ln(2)
        elif kind == "quote":
            pdf.set_text_color(*dim)
            quoting[0] = True
            top = pdf.get_y()
            write_line(block["text"], 11, 6.4, indent=6)
            quoting[0] = False
            pdf.set_draw_color(220, 200, 210)
            pdf.set_line_width(0.8)
            pdf.line(pdf.l_margin + 1.5, top, pdf.l_margin + 1.5, pdf.get_y())
            pdf.set_text_color(*ink)
            pdf.ln(2)
        elif kind == "code":
            pdf.set_font("cjk", "", 9.5)
            pdf.set_fill_color(246, 242, 244)
            pdf.multi_cell(width, 5.2, block["text"] or " ", fill=True,
                           new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(2.5)
        elif kind == "rule":
            pdf.set_draw_color(220, 214, 217)
            pdf.set_line_width(0.3)
            y = pdf.get_y() + 2
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(5)
        elif kind == "table":
            pdf.set_font("cjk", "", 10)
            pdf.set_draw_color(220, 214, 217)
            with pdf.table(width=width, line_height=6) as table:   # 第一行当表头，粗体
                for row in block["rows"]:
                    cells = table.row()
                    for cell in row:
                        cells.cell(plain(cell))
            pdf.set_font("cjk", "", 11)
            pdf.ln(3)
    return bytes(pdf.output())


# ---------------------------------------------------------------- Word 的阅读视图

_PREVIEW_CSS = """
:root { color-scheme: light dark; --ink:#2b2629; --dim:#8a7f84; --line:#e6dde1; --bg:#fbf8f9; --code:#f3edf0; --accent:#c9778b; }
@media (prefers-color-scheme: dark) { :root { --ink:#f1e9ed; --dim:#a99ca2; --line:#4a4045; --bg:#211d1f; --code:#2e282b; --accent:#d896a6; } }
* { box-sizing: border-box; }
body { margin: 0; background: var(--bg); color: var(--ink);
  font: 16px/1.75 -apple-system, "PingFang SC", "Noto Sans SC", sans-serif; -webkit-text-size-adjust: 100%; }
.bar { position: sticky; top: 0; display: flex; align-items: center; gap: 10px; padding: 12px 16px;
  background: color-mix(in srgb, var(--bg) 88%, transparent); -webkit-backdrop-filter: blur(14px); backdrop-filter: blur(14px);
  box-shadow: inset 0 -.5px 0 var(--line); }
.bar a.back { flex: none; width: 34px; height: 34px; margin-left: -6px; border-radius: 50%;
  display: flex; align-items: center; justify-content: center; color: var(--ink); background: none;
  font-size: 28px; font-weight: 400; line-height: 1; padding: 0 2px 4px 0; }
.bar a.back:active { background: var(--code); }
.bar b { flex: 1; min-width: 0; font-size: 14px; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
.dlnote { margin: 12px 16px 0; padding: 10px 14px; border-radius: 12px; background: var(--code);
  color: var(--ink); font-size: 13.5px; line-height: 1.55; }
.dlnote[hidden] { display: none; }
.bar a { flex: none; padding: 7px 14px; border-radius: 999px; background: var(--accent); color: #fff;
  font-size: 13px; font-weight: 600; text-decoration: none; }
main { max-width: 720px; margin: 0 auto; padding: 18px 20px 60px; }
h1, h2, h3, h4 { line-height: 1.35; margin: 1.3em 0 .5em; } h1 { font-size: 1.6em; } h2 { font-size: 1.3em; } h3 { font-size: 1.12em; }
main > :first-child { margin-top: .4em; }
p { margin: .55em 0; } ul, ol { padding-left: 1.4em; margin: .5em 0; } li { margin: .2em 0; }
blockquote { margin: .8em 0; padding: .1em 0 .1em 1em; color: var(--dim); box-shadow: inset 2px 0 0 var(--line); }
pre { white-space: pre-wrap; background: var(--code); padding: 10px 12px; border-radius: 10px; font: 13px/1.6 Menlo, monospace; }
code { background: var(--code); padding: 0 .3em; border-radius: 4px; font-size: .9em; }
hr { border: 0; height: .5px; background: var(--line); margin: 1.4em 0; }
.table { overflow-x: auto; } table { border-collapse: collapse; margin: .8em 0; font-size: 14px; }
td, th { border: .5px solid var(--line); padding: 6px 10px; text-align: left; vertical-align: top; } th { font-weight: 600; }
td.num { text-align: right; font-variant-numeric: tabular-nums; }
.sheet-table td, .sheet-table th { white-space: nowrap; }
.sheet-table td.long { white-space: normal; min-width: 14em; }
.sheet-table td.formula { color: var(--dim); font: 13px Menlo, monospace; }
h2.sheet { font-size: 1.05em; color: var(--dim); margin: 1.6em 0 .4em; }
"""

# 页面上唯一一段脚本，CSP 按它的指纹放行，别的一律不许跑：
# · 返回键：从 dwell 里点进来的就退回去（聊天停在原来的地方）；别处开的就回 dwell 首页。
# · 下载键：加到主屏幕当 App 用时 iPhone 下载不了（新窗口一片空白），点了只出一句提示，
#   让用户去 Safari 里开 dwell 再下；在浏览器里照常下载。
_BACK_SCRIPT = (
    "document.getElementById('back').addEventListener('click',function(e){"
    "var r=document.referrer;"
    "if(r&&r.indexOf(location.origin+'/')===0){e.preventDefault();history.back();}});"
    "document.getElementById('dl').addEventListener('click',function(e){"
    "if(navigator.standalone===true||matchMedia('(display-mode: standalone)').matches){"
    "e.preventDefault();document.getElementById('dlnote').hidden=false;}});"
)
PREVIEW_CSP = (
    "default-src 'none'; style-src 'unsafe-inline'; img-src data:; script-src 'sha256-"
    + base64.b64encode(hashlib.sha256(_BACK_SCRIPT.encode()).digest()).decode() + "'"
)


def _preview_page(title: str, body: str, download_url: str, download_label: str) -> str:
    name = html.escape(title)
    return (
        '<!doctype html><html lang="zh"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>{name}</title><style>{_PREVIEW_CSS}</style></head><body>"
        '<div class="bar"><a id="back" class="back" href="../../../" aria-label="返回">‹</a>'
        # 下载在新窗口开：加到主屏幕当 App 用时，同一个窗口里下载会被系统的「文件」页顶掉，
        # 那页没有返回键；新窗口是 iPhone 自带的小浏览器，左上角有关闭，关了就回到这里
        f'<b>{name}</b><a id="dl" href="{html.escape(download_url)}" target="_blank" rel="noopener">{download_label}</a></div>'
        '<div id="dlnote" class="dlnote" hidden>主屏幕上的 dwell 下载不了文件。在 Safari 里打开 dwell，再点下载。</div>'
        f"<main>{body}</main><script>{_BACK_SCRIPT}</script></body></html>"
    )


def _docx_runs(paragraph) -> str:
    out = []
    for run in paragraph.runs:
        text = html.escape(run.text)
        if not text:
            continue
        mono = (run.font.name or "") in {"Menlo", "Courier New", "Consolas"}
        if mono:
            text = "<code>" + text + "</code>"
        if run.italic:
            text = "<i>" + text + "</i>"
        if run.bold:
            text = "<b>" + text + "</b>"
        out.append(text)
    return "".join(out).replace("\n", "<br>")


def docx_to_html(data: bytes, title: str, download_url: str) -> str:
    """把 Word 读回来排成一页网页：标题、段落、列表、引用、代码、表格。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    open_list = ""

    def close_list():
        nonlocal open_list
        if open_list:
            parts.append("</" + open_list + ">")
            open_list = ""

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            close_list()
            table = Table(child, doc)
            rows = []
            for r, row in enumerate(table.rows):
                cell_tag = "th" if r == 0 else "td"
                cells = "".join(
                    f"<{cell_tag}>" + "<br>".join(_docx_runs(p) for p in cell.paragraphs) + f"</{cell_tag}>"
                    for cell in row.cells
                )
                rows.append("<tr>" + cells + "</tr>")
            parts.append('<div class="table"><table>' + "".join(rows) + "</table></div>")
            continue
        if tag != "p":
            continue
        paragraph = Paragraph(child, doc)
        style = paragraph.style.name if paragraph.style is not None else ""
        body = _docx_runs(paragraph)
        if style.startswith("List"):
            kind = "ol" if "Number" in style else "ul"
            if open_list != kind:
                close_list()
                parts.append("<" + kind + ">")
                open_list = kind
            indent = ' style="margin-left:1.2em"' if style.rstrip().endswith(("2", "3")) else ""
            parts.append(f"<li{indent}>" + body + "</li>")
            continue
        close_list()
        if not body.strip():
            continue
        if style == "Title":
            parts.append("<h1>" + body + "</h1>")
        elif style.startswith("Heading"):
            level = style.split()[-1]
            level = level if level.isdigit() and 1 <= int(level) <= 4 else "4"
            parts.append(f"<h{level}>" + body + f"</h{level}>")
        elif style in {"Quote", "Intense Quote"}:
            parts.append("<blockquote>" + body + "</blockquote>")
        elif paragraph.text.strip() and set(paragraph.text.strip()) <= {"—", "-"}:
            parts.append("<hr>")
        elif body.startswith("<code>") and body.endswith("</code>") and body.count("<code>") == 1:
            parts.append("<pre>" + body[6:-7] + "</pre>")
        else:
            parts.append("<p>" + body + "</p>")
    close_list()
    return _preview_page(title, "".join(parts), download_url, "下载 Word")


# ---------------------------------------------------------------- Excel

_NUMBER = re.compile(r"^[-+]?(\d+(\.\d*)?|\.\d+)([eE][-+]?\d+)?$")
_PERCENT = re.compile(r"^[-+]?(\d+(\.\d*)?|\.\d+)%$")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _cell_value(raw: str):
    """表格里的字变成单元格的值：数字、百分比、公式，其余原样是字。"""
    text = raw.strip()
    if text.startswith("="):
        return text, None
    if _NUMBER.match(text) and not (len(text) > 1 and text[0] == "0" and text[1].isdigit()):
        # 0 开头的长串（电话、编号）留作文字，别把前导 0 吃掉
        number = float(text)
        return (int(number) if number.is_integer() and "." not in text and "e" not in text.lower() else number), None
    if _PERCENT.match(text):
        percent = float(text[:-1])
        # 「0.##%」碰到整数会显示成「80.%」，整数和小数分开给格式
        return percent / 100, ("0%" if percent.is_integer() else "0.0#%")
    return text, None


def _sheet_tables(text: str) -> list[tuple[str, list[list[str]]]]:
    """拆出几张表：Markdown 表格优先（前面的标题当表名）；一张都没有就当 CSV 读。"""
    sheets: list[tuple[str, list[list[str]]]] = []
    title = ""
    for block in blocks(text):
        if block["type"] == "heading":
            title = plain(block["text"])
        elif block["type"] == "table":
            sheets.append((title, [[plain(cell) for cell in row] for row in block["rows"]]))
            title = ""
    if sheets:
        return sheets
    lines = [line for line in text.replace("\r\n", "\n").split("\n") if line.strip()]
    if not lines:
        raise ValueError("Excel 的内容要写成表格（Markdown 表格或 CSV）")
    delimiter = "\t" if "\t" in lines[0] else ","
    rows = [row for row in csv.reader(lines, delimiter=delimiter)]
    return [("", rows)]


def _sheet_name(raw: str, index: int, used: set[str]) -> str:
    name = _BAD_SHEET_CHARS.sub(" ", raw).strip().strip("'")[:31] or f"Sheet{index}"
    base, n = name, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _display_width(value) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(value))


def content_to_xlsx(text: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    header_fill = PatternFill("solid", fgColor="F6ECEF")
    for index, (title, rows) in enumerate(_sheet_tables(text), 1):
        sheet = workbook.create_sheet(_sheet_name(title, index, used))
        widths: dict[int, int] = {}
        for r, row in enumerate(rows, 1):
            for c, raw in enumerate(row, 1):
                value, number_format = _cell_value(str(raw)) if r > 1 else (str(raw).strip(), None)
                cell = sheet.cell(row=r, column=c, value=value)
                if number_format:
                    cell.number_format = number_format
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                cell.alignment = Alignment(vertical="top", wrap_text=_display_width(value) > 50)
                widths[c] = max(widths.get(c, 0), min(_display_width(value), 50))
        for c, width in widths.items():
            sheet.column_dimensions[get_column_letter(c)].width = max(8, width + 2)
        if len(rows) > 1:
            sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xlsx_to_html(data: bytes, title: str, download_url: str) -> str:
    """把 Excel 读回来排成网页：每张工作表一张表，第一行当表头。公式显示公式本身。"""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True)
    parts: list[str] = []
    many = len(workbook.worksheets) > 1
    for sheet in workbook.worksheets:
        if many:
            parts.append('<h2 class="sheet">' + html.escape(sheet.title) + "</h2>")
        rows = []
        for r, row in enumerate(sheet.iter_rows(max_row=2000, values_only=False)):
            cells = []
            for cell in row:
                value = cell.value
                tag = "th" if r == 0 else "td"
                classes = []
                if r and isinstance(value, (int, float)):
                    classes.append("num")
                if isinstance(value, str) and value.startswith("="):
                    classes.append("formula")
                elif isinstance(value, str) and len(value) > 18:
                    classes.append("long")
                cls = f' class="{" ".join(classes)}"' if classes else ""
                if value is None:
                    shown = ""
                elif isinstance(value, (int, float)) and "%" in (cell.number_format or ""):
                    shown = f"{value * 100:.4g}%"
                elif isinstance(value, float):
                    shown = f"{value:.10g}"
                else:
                    shown = str(value)
                cells.append(f"<{tag}{cls}>" + html.escape(shown) + f"</{tag}>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        parts.append('<div class="table"><table class="sheet-table">' + "".join(rows) + "</table></div>")
    workbook.close()
    return _preview_page(title, "".join(parts), download_url, "下载 Excel")


if __name__ == "__main__" and "--ensure-font" in sys.argv:
    regular, bold = ensure_cjk_font()
    print(regular or "（没有中文字体，也没下载到）")
    sys.exit(0 if regular else 1)
